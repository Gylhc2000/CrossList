"""Amazon Flat File（.xlsx）生成 —— openpyxl

对应方案「格式生成：openpyxl（Amazon xlsx）」。
Sheet1 = FlatFile（表头 + 数据行），Sheet2 = 字段说明。
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.rules.platforms import AMAZON_COLUMNS, AMAZON_FIELD_NOTES

_HEADER_FILL = PatternFill("solid", fgColor="FFF3E0")
_HEADER_FONT = Font(bold=True, color="C25A00")


def build_amazon_flatfile(
    data_row: dict[str, object],
    columns: list[str] | None = None,
    notes: dict[str, str] | None = None,
) -> bytes:
    cols = columns or AMAZON_COLUMNS
    note_map = notes or AMAZON_FIELD_NOTES

    wb = Workbook()

    # ---- Sheet1：FlatFile ----
    ws = wb.active
    ws.title = "FlatFile"
    ws.append(cols)
    ws.append([_cell_value(data_row.get(c, "")) for c in cols])

    for idx in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # 列宽：取字段名与内容长度的较大者，上限 60
    for i, c in enumerate(cols, start=1):
        content = str(data_row.get(c, "") or "")
        width = min(max(len(str(c)) + 2, min(len(content) + 2, 60)), 60)
        ws.column_dimensions[get_column_letter(i)].width = width
    for i in range(1, len(cols) + 1):
        ws.cell(row=2, column=i).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    # ---- Sheet2：字段说明 ----
    ws2 = wb.create_sheet("字段说明")
    ws2.append(["字段名", "说明", "当前填充值"])
    for c in cols:
        ws2.append([c, note_map.get(c, ""), str(data_row.get(c, "") or "")])
    for idx in range(1, 4):
        cell = ws2.cell(row=1, column=idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for col, width in zip("ABC", (28, 34, 52)):
        ws2.column_dimensions[col].width = width
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_value(v: object):
    """openpyxl 不接受某些类型，统一转换"""
    if v is None:
        return ""
    if isinstance(v, (str, int, float, bool)):
        return v
    return str(v)
