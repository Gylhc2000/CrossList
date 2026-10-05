"""上架对照表（.xlsx）—— 把批量模板那一行转置成「一字段一行」

为什么要单独出这份表：批量上传模板只对 Amazon Professional 卖家、且只对
"类目模板能下载到手"的那部分人有用；而大多数小卖家是在网页后台一个个表单
填上架的。对他们来说，一列一列的 CSV 没法直接抄，一字段一行、能整格复制的
表才是可用的交付物。图片也同理：平台后台是拖拽上传，需要的不是 URL 而是
「哪个槽位放哪张图」。

纯 CPU 阻塞（openpyxl），调用方须用 asyncio.to_thread 包住。
"""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HDR_FILL = PatternFill("solid", fgColor="FFF3E0")
_HDR_FONT = Font(bold=True, color="C25A00")
_BLOCK_FILL = PatternFill("solid", fgColor="F2F4F7")
_BLOCK_FONT = Font(bold=True, color="3B4453")
_TODO_FONT = Font(color="B54708")
_OK_FONT = Font(color="1B7F4B")

# 状态列文字 → 是否按「待卖家处理」着色。用前缀匹配，状态文案可带具体动作说明。
_TODO_PREFIX = ("待补", "需卖家", "需平台", "先传")

HEADERS = ["字段", "内容（整格复制到平台后台）", "状态", "计量", "说明"]
WIDTHS = (26, 68, 22, 20, 46)


def build_worksheet(title: str, meta: list[tuple[str, str]],
                    blocks: list[tuple[str, list[list[str]]]]) -> bytes:
    """meta: 表头信息（键值对）；blocks: (区块标题, 该区块的行) 列表，行按 HEADERS 对齐。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "上架对照表"

    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13, color="C25A00")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    for k, v in meta:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="6B7688")
        ws.merge_cells(start_row=ws.max_row, start_column=2,
                       end_row=ws.max_row, end_column=len(HEADERS))

    ws.append([])
    for block_title, rows in blocks:
        ws.append([block_title])
        cell = ws.cell(row=ws.max_row, column=1)
        cell.font = _BLOCK_FONT
        for i in range(1, len(HEADERS) + 1):
            ws.cell(row=ws.max_row, column=i).fill = _BLOCK_FILL
        hdr_row = ws.max_row + 1
        ws.append(HEADERS)
        for i in range(1, len(HEADERS) + 1):
            c = ws.cell(row=hdr_row, column=i)
            c.fill = _HDR_FILL
            c.font = _HDR_FONT
        for r in rows:
            ws.append([(x if x is not None else "") for x in r[:len(HEADERS)]])
            status = str(r[2] if len(r) > 2 else "")
            if status.startswith(_TODO_PREFIX):
                ws.cell(row=ws.max_row, column=3).font = _TODO_FONT
            elif status:
                ws.cell(row=ws.max_row, column=3).font = _OK_FONT
        ws.append([])

    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
